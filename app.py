import streamlit as st
import calendar
import os
import json
import pandas as pd
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ReportLab per i PDF
from reportlab.lib.pagesizes import A3, A4, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

st.set_page_config(page_title="Gestione Turni Web", page_icon="📅", layout="wide")

ITALIAN_MONTHS = [
    "Gennaio", "Febbraio", "Marzo", "Aprile", "Maggio", "Giugno",
    "Luglio", "Agosto", "Settembre", "Ottobre", "Novembre", "Dicembre"
]
ITALIAN_DAYS = ["LUN", "MAR", "MER", "GIO", "VEN", "SAB", "DOM"]

ALL_HOURS = ["--:--"]
for h in range(24):
    for m in [0, 30]:
        ALL_HOURS.append(f"{h:02d}:{m:02d}")

def time_to_hours(time_str):
    if not time_str or ":" not in time_str or time_str == "--:--":
        return 0.0
    try:
        parts = time_str.split(":")
        return int(parts[0]) + int(parts[1]) / 60.0
    except Exception:
        return 0.0

def calculate_shift_hours(in_time, out_time):
    in_h = time_to_hours(in_time)
    out_h = time_to_hours(out_time)
    diff = out_h - in_h
    if diff < 0:
        diff += 24.0
    return diff if (in_time and out_time and in_time != "--:--" and out_time != "--:--") else 0.0

# --- INIZIALIZZAZIONE DATO / STATE ---
if "operators" not in st.session_state:
    st.session_state.operators = []
if "clients" not in st.session_state:
    st.session_state.clients = []
if "schedule_data" not in st.session_state:
    st.session_state.schedule_data = {}

# Caricamento file locali se presenti
SAVE_DIR = "SALVATAGGI"
if not os.path.exists(SAVE_DIR):
    os.makedirs(SAVE_DIR)
ANAG_FILE = os.path.join(SAVE_DIR, "anagrafica.json")
SCHED_FILE = os.path.join(SAVE_DIR, "turni_salvati.json")

def load_data_from_disk():
    if os.path.exists(ANAG_FILE):
        try:
            with open(ANAG_FILE, "r", encoding="utf-8") as f:
                data = json.load(f)
                st.session_state.operators = data.get("operators", [])
                st.session_state.clients = data.get("clients", [])
        except Exception:
            pass

    if os.path.exists(SCHED_FILE):
        try:
            with open(SCHED_FILE, "r", encoding="utf-8") as f:
                raw_data = json.load(f)
                st.session_state.schedule_data = {}
                for k, v in raw_data.items():
                    parts = k.split("|")
                    if len(parts) == 5:
                        client, y, m, d, off = parts[0], int(parts[1]), int(parts[2]), int(parts[3]), int(parts[4])
                        st.session_state.schedule_data[(client, y, m, d, off)] = v
        except Exception:
            pass

def save_data_to_disk():
    try:
        with open(ANAG_FILE, "w", encoding="utf-8") as f:
            json.dump({"operators": st.session_state.operators, "clients": st.session_state.clients}, f, ensure_ascii=False, indent=4)
        
        raw_data = {}
        for k, v in st.session_state.schedule_data.items():
            client, y, m, d, off = k
            raw_data[f"{client}|{y}|{m}|{d}|{off}"] = v
        with open(SCHED_FILE, "w", encoding="utf-8") as f:
            json.dump(raw_data, f, ensure_ascii=False, indent=4)
    except Exception as e:
        st.error(f"Errore di salvataggio: {e}")

if not st.session_state.operators and not st.session_state.clients:
    load_data_from_disk()

# --- FUNZIONI DI ESPORTAZIONE (EXCEL & PDF) ---
def generate_excel_bytes(year, month_idx):
    month_name = ITALIAN_MONTHS[month_idx - 1]
    wb = openpyxl.Workbook()
    ws_data = wb.active
    ws_data.title = "Anagrafica"
    ws_sched = wb.create_sheet(title="Pianificazione Mensile")
    ws_res_cl = wb.create_sheet(title="Resoconto Clienti")
    ws_res_op = wb.create_sheet(title="Resoconto Operatori")

    for ws in [ws_data, ws_sched, ws_res_cl, ws_res_op]:
        ws.views.sheetView[0].showGridLines = True

    primary_fill = PatternFill(start_color="366092", fill_type="solid")
    sub_header_fill = PatternFill(start_color="DCE6F1", fill_type="solid")
    accent_fill = PatternFill(start_color="4F81BD", fill_type="solid")
    tot_fill = PatternFill(start_color="F2F2F2", fill_type="solid")
    weekend_fill = PatternFill(start_color="E9EDF4", fill_type="solid")

    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    sub_header_font = Font(name="Calibri", size=9, bold=True, color="1F497D")
    title_font = Font(name="Calibri", size=14, bold=True, color="1F497D")
    section_font = Font(name="Calibri", size=12, bold=True, color="1F497D")
    bold_font = Font(name="Calibri", size=10, bold=True)
    regular_font = Font(name="Calibri", size=10)

    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    side_thin = Side(style='thin', color='B0C4DE')
    side_medium = Side(style='medium', color='366092')
    thin_border = Border(left=side_thin, right=side_thin, top=side_thin, bottom=side_thin)

    ws_data['A1'] = "Gestione Anagrafica Operatori e Clienti"
    ws_data['A1'].font = title_font

    ws_data['A3'] = "ELENCO OPERATORI"
    ws_data['A3'].font = section_font
    ws_data['A4'] = "Operatore"
    ws_data['A4'].fill = primary_fill
    ws_data['A4'].font = header_font

    for i, op in enumerate(st.session_state.operators):
        c = ws_data.cell(row=5+i, column=1, value=op)
        c.font = regular_font; c.border = thin_border

    ws_data['C3'] = "ELENCO CLIENTI"
    ws_data['C3'].font = section_font
    ws_data['C4'] = "Cliente"
    ws_data['C4'].fill = primary_fill
    ws_data['C4'].font = header_font

    for i, cl in enumerate(st.session_state.clients):
        c = ws_data.cell(row=5+i, column=3, value=cl)
        c.font = regular_font; c.border = thin_border

    num_days = calendar.monthrange(year, month_idx)[1]
    ws_sched['A1'] = f"Programmazione Orari - {month_name} {year}"
    ws_sched['A1'].font = title_font
    ws_sched.cell(row=3, column=1, value="Cliente").fill = primary_fill
    ws_sched.cell(row=3, column=1).font = header_font

    total_col = 2 + (num_days * 3)
    ws_sched.cell(row=3, column=total_col, value="Totale Ore Mensili Cliente").fill = primary_fill
    ws_sched.cell(row=3, column=total_col).font = header_font

    for day in range(1, num_days + 1):
        start_col = 2 + (day - 1) * 3
        wday_idx = calendar.weekday(year, month_idx, day)
        wday_str = ITALIAN_DAYS[wday_idx]
        is_weekend = (wday_idx >= 5)
        
        ws_sched.cell(row=3, column=start_col, value=f"{day} {month_name[:3]} ({wday_str})").fill = primary_fill
        ws_sched.cell(row=3, column=start_col).font = header_font
        
        for c_i, h_title in enumerate(["Operatore", "Entrata", "Uscita"]):
            c = ws_sched.cell(row=4, column=start_col+c_i, value=h_title)
            c.fill = weekend_fill if is_weekend else sub_header_fill
            c.font = sub_header_font
        ws_sched.merge_cells(start_row=3, start_column=start_col, end_row=3, end_column=start_col+2)

    for client_idx, client_name in enumerate(st.session_state.clients):
        client_start_row